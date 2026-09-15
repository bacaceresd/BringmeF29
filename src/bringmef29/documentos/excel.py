"""Exporta el Formulario 29 a Excel, con el mismo aspecto que el documento.

Un XLSX sirve para lo que un PDF no: pegar las cifras en una planilla, filtrar
por sección, comparar dos períodos. Por eso el monto va como número —no como
texto con puntos— y cada sección queda separada por una fila en blanco, igual
que en el papel.
"""

from __future__ import annotations

import logging
from pathlib import Path

from ..formulario import Formulario

_log = logging.getLogger(__name__)

# Los mismos colores del documento, en el formato que pide openpyxl.
GRIS_FONDO = "F2F2F7"
GRIS_LINEA = "D9D9DE"
TINTA_SUAVE = "6C6C70"
AZUL_CODIGO = "0A58B0"
AZUL_FONDO = "EEF4FD"
NARANJA = "B25000"
NARANJA_FONDO = "FFF3E0"
VERDE = "1F7A45"
VERDE_FONDO = "E9F7EE"

FORMATO_MONTO = "#,##0;-#,##0;\"\""


class ErrorExcel(RuntimeError):
    """No se pudo generar el archivo."""


def exportar(formulario: Formulario, destino: Path, *, estudio: str = "", vencimiento: str = "") -> str:
    """Escribe el formulario en un XLSX y devuelve la ruta."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as exc:  # pragma: no cover - dependencia opcional
        raise ErrorExcel(
            "Exportar a Excel necesita openpyxl. Instálalo con:  pip install openpyxl"
        ) from exc

    declaracion = formulario.declaracion
    libro = Workbook()
    hoja = libro.active
    hoja.title = f"F29 {declaracion.periodo.codigo}"

    linea_fina = Side(style="thin", color=GRIS_LINEA)
    borde = Border(left=linea_fina, right=linea_fina, top=linea_fina, bottom=linea_fina)
    derecha = Alignment(horizontal="right", vertical="center")
    izquierda = Alignment(horizontal="left", vertical="center", wrap_text=True)
    centro = Alignment(horizontal="center", vertical="center")

    anchos = {"A": 7, "B": 62, "C": 8, "D": 14, "E": 8, "F": 16, "G": 4}
    for columna, ancho in anchos.items():
        hoja.column_dimensions[columna].width = ancho

    fila = 1

    # -- encabezado -----------------------------------------------------
    hoja.cell(fila, 1, "FORMULARIO 29").font = Font(bold=True, size=15)
    hoja.cell(fila, 6, declaracion.periodo.etiqueta).font = Font(bold=True, size=12)
    hoja.cell(fila, 6).alignment = derecha
    fila += 2

    for etiqueta, valor in (
        ("Contribuyente", declaracion.razon_social or "—"),
        ("RUT", declaracion.rut.formateado),
        ("Período tributario", declaracion.periodo.codigo),
        ("Folio", declaracion.folio or "—"),
        ("Formulario", declaracion.procedencia_glosa),
    ):
        hoja.cell(fila, 1, etiqueta).font = Font(size=9, color=TINTA_SUAVE)
        celda = hoja.cell(fila, 2, valor)
        celda.font = Font(bold=True, size=10)
        celda.alignment = izquierda
        fila += 1
    fila += 1

    # -- secciones ------------------------------------------------------
    for seccion in formulario.secciones:
        titulo = seccion.titulo + (f" · {seccion.subtitulo}" if seccion.subtitulo else "")
        celda = hoja.cell(fila, 1, titulo)
        celda.font = Font(bold=True, size=10)
        fila += 1

        cabeceras = ["Línea", "Concepto", "Cód.",
                     seccion.columnas[0] if seccion.columnas else "Cantidad",
                     "Cód.",
                     seccion.columnas[1] if len(seccion.columnas) > 1 else "Monto",
                     ""]
        for columna, texto in enumerate(cabeceras, start=1):
            celda = hoja.cell(fila, columna, texto)
            celda.font = Font(bold=True, size=8, color=TINTA_SUAVE)
            celda.fill = PatternFill("solid", fgColor=GRIS_FONDO)
            celda.border = borde
            celda.alignment = centro if columna in (1, 3, 5, 7) else izquierda
        fila += 1

        for linea in seccion.lineas:
            hoja.cell(fila, 1, linea.numero).alignment = centro
            celda_glosa = hoja.cell(fila, 2, linea.glosa)
            celda_glosa.alignment = izquierda

            if linea.cantidad:
                hoja.cell(fila, 3, linea.cantidad.codigo).alignment = centro
                hoja.cell(fila, 3).font = Font(size=8, bold=True, color=AZUL_CODIGO)
                hoja.cell(fila, 3).fill = PatternFill("solid", fgColor=AZUL_FONDO)
                if linea.cantidad.tiene_valor:
                    valor = hoja.cell(fila, 4, float(linea.cantidad.valor))
                    valor.number_format = FORMATO_MONTO
                    valor.alignment = derecha

            if linea.monto:
                hoja.cell(fila, 5, linea.monto.codigo).alignment = centro
                hoja.cell(fila, 5).font = Font(size=8, bold=True, color=AZUL_CODIGO)
                hoja.cell(fila, 5).fill = PatternFill("solid", fgColor=AZUL_FONDO)
                if linea.monto.tiene_valor:
                    valor = hoja.cell(fila, 6, float(linea.monto.valor))
                    valor.number_format = FORMATO_MONTO
                    valor.alignment = derecha

            hoja.cell(fila, 7, linea.signo).alignment = centro

            for columna in range(1, 8):
                hoja.cell(fila, columna).border = borde
                if linea.total:
                    actual = hoja.cell(fila, columna).font
                    hoja.cell(fila, columna).font = Font(
                        bold=True, size=actual.size, color=actual.color
                    )
            fila += 1

            # Las casillas auxiliares (bases, tasas, créditos) van bajo su línea.
            for casilla in linea.extra:
                if not casilla.tiene_valor and not formulario.completo:
                    continue
                hoja.cell(fila, 2, f"    {casilla.etiqueta}".rstrip()).font = Font(
                    size=9, color=TINTA_SUAVE, italic=True
                )
                hoja.cell(fila, 5, casilla.codigo).alignment = centro
                hoja.cell(fila, 5).font = Font(size=8, color=TINTA_SUAVE)
                if casilla.tiene_valor:
                    valor = hoja.cell(fila, 6, float(casilla.valor))
                    valor.number_format = FORMATO_MONTO
                    valor.alignment = derecha
                    valor.font = Font(size=9, color=TINTA_SUAVE)
                for columna in range(1, 8):
                    hoja.cell(fila, columna).border = borde
                fila += 1

        fila += 1   # el aire entre secciones, igual que en el papel

    # -- cierre ---------------------------------------------------------
    hay_pago = declaracion.hay_que_pagar
    celda = hoja.cell(fila, 2, "TOTAL A PAGAR" if hay_pago else "SIN PAGO ASOCIADO")
    celda.font = Font(bold=True, size=11, color=NARANJA if hay_pago else VERDE)
    celda.fill = PatternFill("solid", fgColor=NARANJA_FONDO if hay_pago else VERDE_FONDO)
    total = hoja.cell(fila, 6, float(declaracion.monto_a_pagar))
    total.number_format = FORMATO_MONTO
    total.font = Font(bold=True, size=13, color=NARANJA if hay_pago else VERDE)
    total.fill = PatternFill("solid", fgColor=NARANJA_FONDO if hay_pago else VERDE_FONDO)
    total.alignment = derecha
    fila += 1

    if hay_pago and vencimiento:
        hoja.cell(fila, 2, f"Vence el {vencimiento}").font = Font(size=9)
        fila += 1
    if estudio:
        fila += 1
        hoja.cell(fila, 2, estudio).font = Font(size=8, color=TINTA_SUAVE)

    hoja.freeze_panes = "A8"
    hoja.print_title_rows = "1:7"
    hoja.page_setup.orientation = "portrait"
    hoja.page_setup.fitToWidth = 1

    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    libro.save(destino)
    _log.info("Excel del formulario generado en %s", destino)
    return str(destino)
